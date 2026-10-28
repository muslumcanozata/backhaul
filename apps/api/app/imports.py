import csv
import io
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.geo import locate
from app.models import TRAILER_CATALOG, Load, TenantSettings
from app.places import find_place, fold

ISTANBUL = ZoneInfo("Europe/Istanbul")
REQUIRED = (
    "referans_no",
    "yukleme_il",
    "yukleme_ilce",
    "bosaltma_il",
    "bosaltma_ilce",
    "tonaj",
    "dorse_tipi",
    "hazir_zamani",
    "teslim_baslangic",
    "teslim_bitis",
)
LOCKED = {"atandi", "yolda", "teslim"}
ROW_LIMIT = 2000


class ImportFileError(Exception):
    def __init__(self, code: str) -> None:
        self.code = code


def read_table(filename: str, raw: bytes) -> list[tuple[int, dict[str, str]]]:
    name = filename.lower()
    headers, rows = _xlsx(raw) if name.endswith(".xlsx") else _csv(raw)
    if not set(REQUIRED).issubset(set(headers)):
        raise ImportFileError("FILE_COLUMNS")
    if len(rows) > ROW_LIMIT:
        raise ImportFileError("FILE_TOO_LARGE")
    return rows


def import_rows(
    db: Session,
    tenant_id,
    settings: TenantSettings,
    rows: list[tuple[int, dict[str, str]]],
    api_key: str | None,
) -> dict:
    accepted = 0
    updated = 0
    rejected: list[dict] = []
    for number, row in rows:
        code, created = _one(db, tenant_id, settings, row, api_key)
        if code:
            rejected.append({"row": number, "reference": (row.get("referans_no") or "").strip(), "code": code})
            continue
        if created:
            accepted += 1
        else:
            updated += 1
    db.commit()
    return {"accepted": accepted, "updated": updated, "rejected": rejected}


def _one(db: Session, tenant_id, settings: TenantSettings, row: dict[str, str], api_key: str | None) -> tuple[str | None, bool]:
    if any(not (row.get(column) or "").strip() for column in REQUIRED):
        return "ROW_REQUIRED", False
    reference = row["referans_no"].strip()
    direction = fold(row.get("yon") or "") or "gidis"
    if direction not in {"gidis", "donus"}:
        return "ROW_DIRECTION", False
    tonnage = _decimal(row["tonaj"])
    if tonnage is None or tonnage <= 0:
        return "ROW_TONNAGE", False
    trailer = fold(row["dorse_tipi"])
    if trailer not in TRAILER_CATALOG or trailer not in settings.enabled_trailer_types:
        return "ROW_TRAILER", False
    pickup = find_place(row["yukleme_il"], row["yukleme_ilce"])
    drop = find_place(row["bosaltma_il"], row["bosaltma_ilce"])
    if pickup is None or drop is None:
        return "ROW_PLACE", False
    ready = _when(row["hazir_zamani"])
    start = _when(row["teslim_baslangic"])
    end = _when(row["teslim_bitis"])
    if ready is None or start is None or end is None or end < ready or end < start:
        return "ROW_WINDOW", False
    freight = _decimal(row.get("navlun_try") or "")
    freight_text = (row.get("navlun_try") or "").strip()
    if direction == "donus" and (freight is None or freight <= 0):
        return "ROW_FREIGHT", False
    if direction == "gidis" and freight_text and (freight is None or freight < 0):
        return "ROW_FREIGHT", False
    if direction == "gidis" and (freight is None or freight <= 0):
        freight = None
    pickup_point = locate(row["yukleme_il"], row["yukleme_ilce"], row.get("yukleme_adres"), api_key)
    drop_point = locate(row["bosaltma_il"], row["bosaltma_ilce"], row.get("bosaltma_adres"), api_key)
    if pickup_point is None or drop_point is None:
        return "ROW_GEOCODE", False
    existing = db.query(Load).filter(Load.tenant_id == tenant_id, Load.reference == reference).one_or_none()
    if existing is not None and existing.status in LOCKED:
        return "ROW_LOCKED", False
    created = existing is None
    load = existing or Load(tenant_id=tenant_id, reference=reference)
    load.direction = direction
    load.customer = (row.get("musteri") or "").strip() or None
    load.broker_firm = (row.get("araci_firma") or "").strip() or None
    load.pickup_il = pickup.il
    load.pickup_ilce = pickup.ilce
    load.pickup_address = (row.get("yukleme_adres") or "").strip() or None
    load.drop_il = drop.il
    load.drop_ilce = drop.ilce
    load.drop_address = (row.get("bosaltma_adres") or "").strip() or None
    load.pickup_lat, load.pickup_lng, pickup_quality = pickup_point
    load.drop_lat, load.drop_lng, drop_quality = drop_point
    load.geocode_quality = "google" if pickup_quality == "google" and drop_quality == "google" else "yaklasik"
    load.tonnage = tonnage
    load.trailer_type = trailer
    load.ready_at = ready
    load.deliver_start = start
    load.deliver_end = end
    load.freight_try = freight
    load.status = "havuz" if direction == "donus" else "planlanabilir"
    load.vehicle_id = None
    if created:
        db.add(load)
    db.flush()
    return None, created


def _csv(raw: bytes) -> tuple[list[str], list[tuple[int, dict[str, str]]]]:
    text = raw.decode("utf-8-sig")
    first = text.splitlines()[0] if text.splitlines() else ""
    delimiter = ";" if first.count(";") > first.count(",") else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    headers = [(name or "").strip() for name in (reader.fieldnames or [])]
    rows: list[tuple[int, dict[str, str]]] = []
    for index, item in enumerate(reader, start=2):
        cleaned = {(key or "").strip(): (value or "").strip() for key, value in item.items()}
        if any(cleaned.values()):
            rows.append((index, cleaned))
    return headers, rows


def _xlsx(raw: bytes) -> tuple[list[str], list[tuple[int, dict[str, str]]]]:
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheet = workbook.active
    matrix = list(sheet.iter_rows(values_only=True))
    workbook.close()
    if not matrix:
        return [], []
    headers = [str(cell).strip() if cell is not None else "" for cell in matrix[0]]
    rows: list[tuple[int, dict[str, str]]] = []
    for offset, values in enumerate(matrix[1:], start=2):
        record: dict[str, str] = {}
        for key, value in zip(headers, values, strict=False):
            if not key:
                continue
            if isinstance(value, datetime):
                record[key] = value.strftime("%Y-%m-%d %H:%M")
            elif value is None:
                record[key] = ""
            else:
                record[key] = str(value).strip()
        if any(record.values()):
            rows.append((offset, record))
    return headers, rows


EXPORT_HEADERS = (
    "referans_no",
    "yukleme_il",
    "yukleme_ilce",
    "bosaltma_il",
    "bosaltma_ilce",
    "tonaj",
    "dorse_tipi",
    "hazir_zamani",
    "teslim_baslangic",
    "teslim_bitis",
    "yon",
    "navlun_try",
    "musteri",
    "araci_firma",
    "yukleme_adres",
    "bosaltma_adres",
    "durum",
)


def export_workbook(loads: list[Load]) -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "yukler"
    sheet.append(list(EXPORT_HEADERS))
    for load in loads:
        sheet.append(
            [
                load.reference,
                load.pickup_il,
                load.pickup_ilce,
                load.drop_il,
                load.drop_ilce,
                float(load.tonnage),
                load.trailer_type,
                _clock_out(load.ready_at),
                _clock_out(load.deliver_start),
                _clock_out(load.deliver_end),
                load.direction,
                float(load.freight_try) if load.freight_try is not None else "",
                load.customer or "",
                load.broker_firm or "",
                load.pickup_address or "",
                load.drop_address or "",
                load.status,
            ]
        )
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def create_from_form(db: Session, tenant_id, settings: TenantSettings, body: dict, api_key: str | None) -> tuple[str | None, Load | None]:
    row = {
        "referans_no": body["reference"],
        "yukleme_il": body["pickup_il"],
        "yukleme_ilce": body["pickup_ilce"],
        "bosaltma_il": body["drop_il"],
        "bosaltma_ilce": body["drop_ilce"],
        "tonaj": str(body["tonnage"]),
        "dorse_tipi": body["trailer_type"],
        "hazir_zamani": _clock_out(body["ready_at"]),
        "teslim_baslangic": _clock_out(body["deliver_start"]),
        "teslim_bitis": _clock_out(body["deliver_end"]),
        "yon": body.get("direction") or "gidis",
        "navlun_try": "" if body.get("freight_try") is None else str(body["freight_try"]),
        "musteri": body.get("customer") or "",
        "araci_firma": body.get("broker_firm") or "",
        "yukleme_adres": body.get("pickup_address") or "",
        "bosaltma_adres": body.get("drop_address") or "",
    }
    code, created = _one(db, tenant_id, settings, row, api_key)
    if code:
        return code, None
    load = db.query(Load).filter(Load.tenant_id == tenant_id, Load.reference == row["referans_no"]).one()
    return None, load


def _clock_out(value: datetime) -> str:
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    return value.astimezone(ISTANBUL).strftime("%Y-%m-%d %H:%M")


def _decimal(value: str) -> Decimal | None:
    text = value.strip().replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _when(value: str) -> datetime | None:
    text = value.strip()
    for fmt in (
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%dT%H:%M",
        "%Y-%m-%dT%H:%M:%S",
        "%d.%m.%Y %H:%M",
        "%d.%m.%Y %H:%M:%S",
    ):
        try:
            parsed = datetime.strptime(text, fmt)
            return parsed.replace(tzinfo=ISTANBUL).astimezone(UTC)
        except ValueError:
            continue
    return None
